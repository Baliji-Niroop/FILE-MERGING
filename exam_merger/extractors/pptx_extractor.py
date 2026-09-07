import os

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    Presentation = None
    MSO_SHAPE_TYPE = None

from ..models.chunk import Chunk
from ..diagrams.vision_describer import describe_image

def _image_media_type(image_ext: str) -> str:
    ext = (image_ext or "").lower().lstrip('.')
    mapping = {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "gif": "image/gif",
        "bmp": "image/bmp",
        "tif": "image/tiff",
        "tiff": "image/tiff",
        "webp": "image/webp",
    }
    return mapping.get(ext, "image/png")

def pptx_table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        row_cells = []
        for cell in row.cells:
            row_cells.append(cell.text.replace("\n", " ").strip())
        if any(cell != "" for cell in row_cells):
            rows.append(row_cells)
            
    if not rows:
        return ""
        
    cols_count = len(rows[0])
    md_lines = []
    # Header
    md_lines.append("| " + " | ".join(rows[0]) + " |")
    # Separator
    md_lines.append("| " + " | ".join(["---"] * cols_count) + " |")
    # Rows
    for row in rows[1:]:
        if len(row) < cols_count:
            row.extend([""] * (cols_count - len(row)))
        elif len(row) > cols_count:
            row = row[:cols_count]
        md_lines.append("| " + " | ".join(row) + " |")
        
    return "\n".join(md_lines)

def extract(path: str, describe_diagrams: bool = False) -> list[Chunk]:
    chunks = []
    filename = os.path.basename(path)

    if Presentation is None:
        return []
    
    try:
        prs = Presentation(path)
    except Exception as e:
        print(f"Error reading {path} via Presentation: {e}")
        return []
        
    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        location_prefix = f"slide {slide_num}"
        
        # 1. Extract Slide Title
        title_text = ""
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text_frame.text.strip()
        except AttributeError:
            # Some slide layouts may not support .title attributes cleanly or raise exception
            pass
            
        if title_text:
            chunks.append(Chunk(
                text=title_text,
                source_file=filename,
                location=f"{location_prefix} - title",
                chunk_type="text"
            ))
            
        # 2. Extract Other Text Boxes / Shapes
        body_parts = []
        table_idx = 1
        image_idx = 1
        
        for shape in slide.shapes:
            # Skip title shape since we already captured it
            try:
                if slide.shapes.title and shape == slide.shapes.title:
                    continue
            except AttributeError:
                pass
                
            # Text frames
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    body_parts.append(text)
                    
            # Tables
            if shape.has_table:
                table_md = pptx_table_to_markdown(shape.table)
                if table_md:
                    chunks.append(Chunk(
                        text=table_md,
                        source_file=filename,
                        location=f"{location_prefix} - table {table_idx}",
                        chunk_type="table"
                    ))
                    table_idx += 1
                    
            # Embedded pictures (optional vision pass)
            if describe_diagrams and MSO_SHAPE_TYPE is not None and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image_bytes = shape.image.blob
                    description = describe_image(image_bytes, _image_media_type(getattr(shape.image, "ext", "")))
                    if description:
                        chunks.append(Chunk(
                            text=description,
                            source_file=filename,
                            location=f"{location_prefix} - diagram {image_idx}",
                            chunk_type="diagram_description"
                        ))
                        image_idx += 1
                except Exception as e:
                    print(f"Warning: Failed to describe image on slide {slide_num} in {filename}: {e}")
                    
        # Add body texts chunk if we found any
        if body_parts:
            combined_body = "\n\n".join(body_parts)
            chunks.append(Chunk(
                text=combined_body,
                source_file=filename,
                location=location_prefix,
                chunk_type="text"
            ))
            
        # 3. Extract Speaker Notes
        notes_text = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass
            
        if notes_text:
            chunks.append(Chunk(
                text=notes_text,
                source_file=filename,
                location=f"{location_prefix} - notes",
                chunk_type="speaker_notes"
            ))
            
    return chunks
